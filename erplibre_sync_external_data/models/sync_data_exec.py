import base64
import csv
import io
import json
import logging
import re
from collections import defaultdict
from datetime import date, datetime, time

import pytz
from markupsafe import Markup
from odoo import _, api, conf, fields, models
from odoo.exceptions import ValidationError
from odoo.tools.mimetypes import guess_mimetype

try:
    # openpyxl is the safest lib for .xlsx on Odoo 18 (Python >= 3.10)
    import openpyxl
except ImportError:
    openpyxl = None
_logger = logging.getLogger(__name__)


class SyncDataExec(models.Model):
    _name = "sync.data.exec"
    _description = "Synchronisation with Excel"
    _inherit = ["mail.activity.mixin", "mail.thread"]
    _order = "id desc"

    name = fields.Char(compute="_compute_name", tracking=True)

    state = fields.Selection(
        [
            ("upload", "Téléversement"),
            ("summary", "Résumé"),
        ],
        default="upload",
        required=True,
        tracking=True,
    )

    notes = fields.Html(string="Notes", readonly=True, tracking=True)

    sync_model_ids = fields.Many2many(
        "sync.model",
        string="Template sync",
        tracking=True,
    )

    transform_id = fields.Many2one(
        "sync.data.transform",
        string="Transformation",
        readonly=True,
        help="Latest transformation sheet",
        tracking=True,
    )

    what_import = fields.Selection(
        selection=[("default", "Default")],
        string="What to import",
        default="default",
        tracking=True,
    )

    context_name = fields.Selection(
        selection=[], string="Context name", tracking=True
    )

    has_modification = fields.Boolean(tracking=True)

    modif_count = fields.Integer(tracking=True, readonly=True)

    create_count = fields.Integer(tracking=True, readonly=True)

    has_error = fields.Boolean(tracking=True)

    has_error_msg = fields.Text(tracking=True)

    sync_data_write_ids = fields.One2many(
        comodel_name="mail.tracking.value",
        inverse_name="sync_data_exec_id",
        domain=[("field_id.name", "!=", "file_no_line")],
    )

    send_message_at_create_or_modify_sync = fields.Boolean(
        string="Send notification at change", default=True
    )

    send_message_force = fields.Boolean(string="Force Send notification")

    time_execution_extract_start = fields.Datetime(
        readonly=True, tracking=True
    )

    time_execution_extract_end = fields.Datetime(readonly=True, tracking=True)

    time_duration_extract = fields.Float(
        compute="_compute_time_duration_extract", store=True, tracking=True
    )

    time_duration_extract_fr = fields.Char(
        compute="_compute_time_duration_extract", store=True, tracking=True
    )

    sync_data_create_ids = fields.One2many(
        comodel_name="sync.data.create",
        inverse_name="sync_data_exec_id",
        tracking=True,
    )

    @api.model_create_multi
    def create(self, vals_list):
        res = super().create(vals_list)

        res.message_subscribe()

        return res

    def message_subscribe(self, partner_ids=None, subtype_ids=None):
        group_sync_id = self.env.ref(
            "erplibre_sync_external_data.group_erplibre_sync_external_data_exec_notify"
        )
        group_partner_ids = group_sync_id.users.mapped("partner_id").ids
        partner_ids = (partner_ids or []) + group_partner_ids

        return super().message_subscribe(
            partner_ids=partner_ids, subtype_ids=subtype_ids
        )

    def action_send_message_notification(self):
        mail_template = self.env.ref(
            "erplibre_sync_external_data.mail_template_erplibre_sync_external_data_notif_change"
        )
        for rec in self:
            if not (
                rec.send_message_force
                or (
                    rec.send_message_at_create_or_modify_sync
                    and (rec.sync_data_create_ids or rec.sync_data_write_ids)
                )
            ):
                continue

            partners = rec.message_partner_ids
            partners -= self.env.user.partner_id
            partners = partners.filtered(lambda p: p.email)

            if not partners:
                continue
            body_html = mail_template._render_field(
                "body_html", [rec.id], compute_lang=True
            )[rec.id]
            body_html = Markup(body_html)
            subject = mail_template._render_field(
                "subject", rec.ids, compute_lang=True
            )[rec.id]

            rec.message_post(
                body=body_html,
                subject=subject,
                message_type="comment",
                subtype_xmlid="mail.mt_comment",
                partner_ids=partners.ids,
            )

    @api.depends("what_import", "time_execution_extract_end")
    def _compute_name(self):
        for rec in self:
            if rec.what_import:
                selection = dict(self._fields["what_import"].selection)
                selection_name = (
                    selection.get(rec.what_import) or rec.what_import
                )
                if rec.time_execution_extract_end:
                    rec.name = f"{selection_name} - {rec.time_execution_extract_end} UDT"
                else:
                    rec.name = selection_name
            else:
                rec.name = False

    @api.depends("time_execution_extract_start", "time_execution_extract_end")
    def _compute_time_duration_extract(self):
        for rec in self:
            duration, label = rec._compute_duration(
                rec.time_execution_extract_start,
                rec.time_execution_extract_end,
            )
            if (
                not rec.time_execution_extract_start
                and rec.time_execution_extract_end
            ):
                rec.time_execution_extract_start = (
                    rec.time_execution_extract_end
                )
            rec.time_duration_extract = duration
            rec.time_duration_extract_fr = label

    def _compute_duration(self, start, end):
        """Compute duration in seconds and a human-readable label."""
        if not start and not end:
            return 0, _("Not running")
        if start and not end:
            return 0, _("Running")
        if not start or end < start:
            return 0, self._format_duration_fr(0)
        delta = end - start
        seconds = int(delta.total_seconds())
        return float(seconds), self._format_duration_fr(seconds)

    def _format_duration_fr(self, seconds):
        """Return a human-readable duration string from a number of seconds."""
        if not seconds or seconds < 0:
            return _("0 seconde")

        days, rem = divmod(seconds, 86400)
        hours, rem = divmod(rem, 3600)
        minutes, seconds = divmod(rem, 60)

        parts = []

        if days:
            parts.append(_("%d day(s)") % days)
        if hours:
            parts.append(_("%d hour(s)") % hours)
        if minutes:
            parts.append(_("%d minute(s)") % minutes)
        if seconds:
            parts.append(_("%d second(s)") % seconds)

        return ", ".join(parts) if parts else _("0 second")

    def action_process_sync_data_async(self):
        for rec in self:
            rec = rec.exists()
            if not rec:
                continue
            if "queue_job" in conf.server_wide_modules:
                rec.with_delay().action_process_sync_data()
            else:
                rec.action_process_sync_data()

    REVERT_FIELD_MAP = {
        "date": "old_value_datetime",
        "datetime": "old_value_datetime",
        "int": "old_value_integer",
        "float": "old_value_float",
        "monetary": "old_value_float",
        "char": "old_value_char",
        "text": "old_value_text",
    }

    def action_revert_data(self):
        for rec in self:
            for diff_id in rec.sync_data_write_ids:
                old_value_field = self.REVERT_FIELD_MAP.get(
                    diff_id.field_id.ttype
                )
                if old_value_field:
                    rec_id = self.env[diff_id.model].browse(diff_id.res_id)
                    setattr(
                        rec_id,
                        diff_id.field_id.name,
                        getattr(diff_id, old_value_field),
                    )

    def action_process_transform_async(self):
        for rec in self:
            rec = rec.exists()
            if not rec:
                continue
            if "queue_job" in conf.server_wide_modules:
                rec.with_delay().action_process_transform({})
            else:
                rec.action_process_transform({})

    def action_process_sync_data(self):
        self.ensure_one()
        if not self.time_execution_extract_start:
            self.time_execution_extract_start = fields.Datetime.now()

        if self.what_import == "default":
            self.action_import_default_algo()

        return {}

    def action_import_default_algo(self):
        for rec in self:
            if not rec.file:
                raise ValidationError("Missing file")
            # TODO depend from erplibre_sync_external_data_spreadsheet, maybe merge modules
            raw_data = base64.b64decode(rec.file)
            mimetype = guess_mimetype(raw_data)
            # TODO how can retrieve original name
            attachment_id = self.env["ir.attachment"].create(
                {
                    "name": "default_file.xlsx",
                    "datas": rec.file,
                    "res_model": rec._name,
                    "res_id": rec.id,
                    "mimetype": mimetype,
                }
            )
            subtype_note = self.env.ref("mail.mt_note")
            self.env["mail.message"].create(
                [
                    {
                        "model": rec._name,
                        "res_id": rec.id,
                        "body": _("File added automatically."),
                        "message_type": "comment",
                        "subtype_id": subtype_note.id,
                        "author_id": self.env.user.partner_id.id,
                        "attachment_ids": attachment_id.ids,
                    }
                ]
            )

            filepath_byte = io.BytesIO(raw_data)
            for sync_model_id in rec.sync_model_ids:
                rec.extract_automated_excel(
                    filepath_byte,
                    sync_model_id,
                )
            rec.end_time_execution()

    @staticmethod
    def _ensure_openpyxl():
        if openpyxl is None:
            raise ImportError(
                "openpyxl is required to process .xlsx files. "
                "Install it with: pip install openpyxl"
            )

    @staticmethod
    def _extract_cell_value(cell, is_excel):
        """Extract a typed value from a single cell."""
        if not is_excel:
            return cell
        if isinstance(cell.value, float) and cell.number_format == "0":
            return int(cell.value)
        if cell.number_format == "mm-dd-yy" and cell.base_date:
            return cell.base_date
        return cell.value

    @staticmethod
    def _extract_header_values(row, is_excel):
        """Extract cleaned header values from a row."""
        if is_excel:
            values = []
            for cell in row:
                if cell.value is None:
                    values.append("")
                else:
                    values.append(
                        cell.value.strip()
                        .replace("\n", " ")
                        .replace("\t", " ")
                    )
            return values
        return [
            a.lstrip("\ufeff")
            .strip()
            .strip('"')
            .replace("\n", "")
            .replace("\t", "")
            for a in row
        ]

    def _open_file_reader(self, filepath, filetype, sheet_name, index):
        """Open a file and return (reader, is_excel, index) or (None, False, index) on error."""
        if filetype == "xlsx":
            self._ensure_openpyxl()
            wb = openpyxl.load_workbook(filepath, data_only=True)
            if not wb.sheetnames:
                raise ValueError("Missing sheet from xlsx file.")
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb[wb.sheetnames[0]]
            return ws.iter_rows(), True, index
        if filetype == "csv":
            index -= 1
            if isinstance(filepath, str):
                reader = []
                with open(filepath, newline="", encoding="utf-8") as f:
                    for row in csv.reader(f):
                        reader.append(list(row))
            else:
                text_stream = io.TextIOWrapper(filepath, encoding="utf-8")
                reader = csv.reader(text_stream)
            return reader, False, index
        _logger.error("Unsupported filetype '%s'", filetype)
        return None, False, index

    def extract_automated_excel(self, filepath, sync_model_id):
        last_id_tracking = (
            self.env["mail.tracking.value"]
            .search([], order="id desc", limit=1)
            .id
        )
        model_data_to_track = defaultdict(list)
        index = 0
        parsed_headers = []
        data_lines = []
        last_repeat_column_data = []
        has_different_header = False

        file_name_type = sync_model_id.name
        model_name = sync_model_id.model_name

        metadata = json.loads(sync_model_id.spreadsheet_extraction_metadata)
        header_config = metadata.get("header", [])
        header_config_ext = header_config + [("File no line", "file_no_line")]
        expected_headers = [a[0] for a in header_config]
        index_line_header = metadata.get("index_line_header", 0)
        nb_line_header = metadata.get("nb_line_header", 1)
        callback_init_read_header = metadata.get("callback_init_read_header")
        sync_fields = metadata.get("sync", [])
        ignore_last_line = metadata.get("ignore_last_line", 0)
        index_last_line = metadata.get("index_last_line", 0)
        ignore_data = metadata.get("ignore_data", [])
        ignore_validation_header = metadata.get(
            "ignore_validation_header", False
        )
        sheet_name = metadata.get("sheet_name", "")
        filetype = metadata.get("filetype", "")

        option_level = metadata.get("option_level", {})
        # repeat_column will ignore the line when need repetition and will move data to next line
        repeat_column = option_level.get("repeat_column", [])
        # duplicate_column_when_empty will replicate the value when empty, like buffering
        duplicate_columns = option_level.get("duplicate_column_when_empty", [])
        buffer_duplicate_column_when_empty = [None] * len(header_config)

        sync_header_indices = [
            i for i, a in enumerate(header_config_ext) if a[1] in sync_fields
        ]

        reader, is_excel, index = self._open_file_reader(
            filepath, filetype, sheet_name, index
        )
        if reader is None:
            return

        finish_read_header = False
        for row in reader:
            index += 1
            if index_line_header > index or (
                index_last_line > 0 and index > index_last_line
            ):
                continue
            if ignore_validation_header:
                parsed_headers = expected_headers

            is_other_header = (
                index_line_header
                < index
                <= index_line_header + nb_line_header - 1
            )

            if (
                index_line_header
                <= index
                <= index_line_header + nb_line_header - 1
            ):

                row_values = self._extract_header_values(row, is_excel)
                for item_row_i, item_row in enumerate(row_values):
                    item_row_transform = item_row.strip()

                    if len(header_config) <= item_row_i:
                        _logger.warning(
                            "Got difference header, missing index %s '%s', and '%s' "
                            "check filepath '%s' filetype '%s'",
                            item_row_i,
                            header_config,
                            parsed_headers,
                            filepath,
                            file_name_type,
                        )
                        has_different_header = True
                    elif (
                        item_row_transform != header_config[item_row_i][0]
                        and not is_other_header
                    ):
                        _logger.warning(
                            "Got difference header '%s' and '%s', "
                            "check filepath '%s' filetype '%s'",
                            item_row_transform,
                            header_config[item_row_i][0],
                            filepath,
                            file_name_type,
                        )
                        has_different_header = True
                    if is_other_header:
                        # Append if multiple header
                        if len(parsed_headers) > item_row_i:
                            parsed_headers[
                                item_row_i
                            ] += f" {item_row_transform}"
                    else:
                        parsed_headers.append(item_row_transform)
            elif index > index_line_header:
                if not finish_read_header and callback_init_read_header:
                    cb_init_read_header = getattr(
                        self, callback_init_read_header
                    )
                    if cb_init_read_header:
                        parsed_headers = cb_init_read_header(
                            parsed_headers, expected_headers
                        )
                finish_read_header = True
                row_values = []
                max_cell_index = min(len(parsed_headers), len(header_config))
                for index_cell, cell_sheet in enumerate(row):
                    if index_cell >= max_cell_index:
                        break
                    value_mapping = {}
                    if len(header_config[index_cell]) > 2:
                        value_mapping = header_config[index_cell][2]
                    value = self._extract_cell_value(cell_sheet, is_excel)
                    if value in ignore_data:
                        value = False
                    if value_mapping:
                        new_value = value_mapping.get(value)
                        if new_value:
                            value = new_value
                    row_values.append(value)
                if any(
                    row_values[1:]
                ):  # TODO it's not a generic check, or don't use repeat, maybe need from index_sync_header
                    row_values.append(index)
                    if duplicate_columns:
                        for index_duplicate in duplicate_columns:
                            if not (0 <= index_duplicate < len(row_values)):
                                raise ValidationError(
                                    f"Check your configuration, duplicate_column_when_empty {duplicate_columns}, index is wrong with values."
                                )
                            value = row_values[index_duplicate]
                            if value:
                                buffer_duplicate_column_when_empty[
                                    index_duplicate
                                ] = value
                            else:
                                row_values[index_duplicate] = (
                                    buffer_duplicate_column_when_empty[
                                        index_duplicate
                                    ]
                                )
                    # Support repeat column
                    if repeat_column:
                        detect_column_repeat = [
                            row_values[a] for a in repeat_column
                        ]
                        if any(detect_column_repeat):
                            last_repeat_column_data = row_values
                        else:
                            if last_repeat_column_data:
                                for i in repeat_column:
                                    row_values[i] = last_repeat_column_data[i]
                                data_lines.append(row_values)
                            else:
                                _logger.warning(
                                    "Missing information from repeat column option."
                                )
                    else:
                        if sync_fields:
                            # Validate sync value exist
                            sync_check_values = [
                                row_values[a] for a in sync_header_indices
                            ]
                            # TODO support [0], because it's false
                            if any(sync_check_values):
                                data_lines.append(row_values)
                        else:
                            data_lines.append(row_values)

        if has_different_header and not data_lines:
            _logger.error(
                "Wrong header file '%s', expected header '%s' and got '%s'",
                file_name_type,
                expected_headers,
                parsed_headers,
            )
            return
        if ignore_last_line:
            data_lines = data_lines[:-ignore_last_line]
        for line in data_lines:
            record_values, file_no_line = self._parse_data_line(
                line, header_config, model_name
            )
            model_record_id = self._sync_record(
                model_name,
                record_values,
                sync_fields,
                file_no_line,
                file_name_type,
            )
            model_data_to_track[model_name].append(model_record_id.id)

        self._link_tracking_values(model_data_to_track, last_id_tracking)

    def _parse_data_line(self, line, header_config, model_name):
        """Parse a data line into a record dict and file_no_line value."""
        record_values = {}
        file_no_line = -1
        for index_column, column_value in enumerate(line):
            if index_column == len(line) - 1:
                field_name = "file_no_line"
                file_no_line = column_value
            else:
                field_name = header_config[index_column][1]
            field_obj = self.env[model_name]._fields.get(field_name)
            if not field_obj:
                raise ValidationError(
                    f"Field '{field_name}' not found on model"
                    f" '{model_name}'"
                )
            field_type = field_obj.type
            record_values[field_name] = self._convert_field_value(
                record_values, field_name, field_type, column_value
            )
        return record_values, file_no_line

    def _sync_record(
        self,
        model_name,
        record_values,
        sync_fields,
        file_no_line,
        file_name_type,
    ):
        """Find or create a record, updating if it already exists."""
        search_domain = [
            (field, "=", record_values[field]) for field in sync_fields
        ]
        model_record_id = self.env[model_name].search(search_domain)

        if len(model_record_id) > 1:
            _logger.warning(
                "Found multiple records with index %s", sync_fields
            )
            matching = [
                r for r in model_record_id if r.file_no_line == file_no_line
            ]
            if len(matching) != 1:
                raise ValidationError(
                    f"Cannot resolve duplicate, check "
                    f"{[r.id for r in model_record_id]} of {file_name_type}"
                )
            model_record_id = matching[0]

        if not model_record_id:
            model_record_id = self.env[model_name].create([record_values])
            self.env["sync.data.create"].create(
                [
                    {
                        "res_model": model_name,
                        "res_id": model_record_id.id,
                        "sync_data_exec_id": self.id,
                    }
                ]
            )
        else:
            model_record_id.write(record_values)
        return model_record_id

    def _link_tracking_values(self, model_data_to_track, last_id_tracking):
        """Link mail.tracking.value records to this sync execution."""
        if not model_data_to_track:
            return
        self.env.cr.commit()
        for tracked_model, res_ids in model_data_to_track.items():
            tracking_vals = self.env["mail.tracking.value"].search(
                [
                    ("mail_message_id.res_model", "=", tracked_model),
                    ("mail_message_id.res_id", "in", res_ids),
                    ("id", ">", last_id_tracking),
                ]
            )
            if tracking_vals:
                tracking_vals.sync_data_exec_id = self.id

    def action_process_transform(self, sync_data_transform_value: dict = None):
        self.ensure_one()
        self.time_execution_extract_start = fields.Datetime.now()
        if not sync_data_transform_value:
            sync_data_transform_value = {}
        if "context_name" not in sync_data_transform_value:
            sync_data_transform_value["context_name"] = (
                self.context_name or "default"
            )

        sync_data_transform_value["sync_model_ids"] = [
            (6, 0, self.sync_model_ids.ids)
        ]
        self.transform_id = self.env["sync.data.transform"].create(
            [sync_data_transform_value]
        )
        self.transform_id.action_transform_algo()
        return {}

    def end_time_execution(self):
        for rec in self:
            rec.update_modification_value()
            rec.action_send_message_notification()
            rec.write(
                {
                    "time_execution_extract_end": fields.Datetime.now(),
                    "state": "summary",
                }
            )

    def update_modification_value(self):
        for rec in self:
            rec.create_count = len(rec.sync_data_create_ids)
            rec.modif_count = len(rec.sync_data_write_ids)
            rec.has_modification = bool(rec.modif_count or rec.create_count)

    def action_copy(self):
        self.ensure_one()

        copied_record = self.copy()
        copied_record.has_error = False
        copied_record.has_error_msg = False
        copied_record.action_process_sync_data()
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "view_mode": "form",
            "res_id": copied_record.id,
            "target": "current",
        }

    def _convert_field_value(
        self, record_values, field_name, field_type, value
    ):
        """Convert a raw cell value to the appropriate Odoo field type."""
        if field_type in ("date", "datetime"):
            record_values[field_name] = value
            if not isinstance(value, fields.datetime):
                self._transform_date(record_values, field_name)
            return record_values[field_name]
        if field_type == "boolean":
            if isinstance(value, str):
                return value.lower() in ("o", "y", "yes", "oui", "true")
            return value
        if field_type == "integer":
            return self._parse_integer(value)
        if field_type in ("float", "monetary"):
            return self._parse_float(value)
        if field_type in ("char", "text"):
            if not isinstance(value, str):
                return str(value) if value else ""
            return value
        return value

    def _parse_integer(self, value):
        """Parse a string value into an integer, stripping non-numeric chars."""
        if not isinstance(value, str):
            return value
        if not value:
            return 0
        cleaned = re.sub(r"[^\d-]", "", value)
        if value.replace(" ", "") != cleaned:
            _logger.error("Detect int %s from char %s", cleaned, value)
        return int(cleaned) if cleaned else 0

    def _parse_float(self, value):
        """Parse a string value into a float, stripping currency symbols."""
        if not isinstance(value, str):
            return value
        value = value.replace("$", "")
        if not value:
            return 0.0
        cleaned = re.sub(r"[^\d,.-]", "", value).replace(",", ".")
        if value.replace(" ", "") != cleaned:
            _logger.error("Detect float %s from char %s", cleaned, value)
        return float(cleaned) if cleaned else 0.0

    FRENCH_MONTHS = {
        "janv": 1,
        "jan": 1,
        "févr": 2,
        "fevr": 2,
        "fév": 2,
        "mars": 3,
        "avr": 4,
        "mai": 5,
        "juin": 6,
        "juil": 7,
        "jul": 7,
        "août": 8,
        "aout": 8,
        "aoû": 8,
        "sept": 9,
        "oct": 10,
        "nov": 11,
        "déc": 12,
        "dec": 12,
    }

    def _transform_date(self, record_data, key):
        data = record_data.get(key)
        if data is None or isinstance(data, date):
            return
        if isinstance(data, int):
            record_data[key] = datetime.strptime(str(data), "%Y").date()
            return
        if not isinstance(data, str) or not data:
            record_data[key] = False
            return

        sep = "/" if "/" in data else "-" if "-" in data else None
        if not sep:
            record_data[key] = False
            return

        # Try standard formats first
        formats = [
            f"%d{sep}%m{sep}%Y",
            f"%Y{sep}%m{sep}%d",
            f"%d{sep}%b{sep}%y",
            f"%d{sep}%B{sep}%y",
        ]
        for fmt in formats:
            try:
                record_data[key] = datetime.strptime(data, fmt).date()
                return
            except (ValueError, TypeError):
                continue

        # Replace French month abbreviations and retry
        normalized = data.lower()
        for month_abbr, month_num in self.FRENCH_MONTHS.items():
            normalized = normalized.replace(month_abbr, str(month_num))

        short_long_formats = [
            f"%d{sep}%m{sep}%y",
            f"%m{sep}%d{sep}%y",
            f"%m{sep}%d{sep}%Y",
        ]
        for fmt in short_long_formats:
            try:
                record_data[key] = datetime.strptime(normalized, fmt).date()
                return
            except (ValueError, TypeError):
                continue

        record_data[key] = False
        _logger.error("Cannot parse data to date '%s'", data)

    def transform_datetime(self, user_datetime):
        """Convert a date to a UTC datetime string using the user's timezone."""
        if not user_datetime:
            return user_datetime
        tz_name = self.env.user.tz or "America/Toronto"
        tz = pytz.timezone(tz_name)
        dt_local = tz.localize(datetime.combine(user_datetime, time.min))
        dt_utc = dt_local.astimezone(pytz.UTC)
        return fields.Datetime.to_string(dt_utc)
